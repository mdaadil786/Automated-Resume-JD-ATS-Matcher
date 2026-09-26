import io
import os
import zipfile
from pathlib import Path
import pdfplumber
from docx import Document
from app.config import settings

ALLOWED = {".pdf", ".docx", ".txt", ".zip"}

def safe_name(name):
    return os.path.basename(name or "uploaded_file").replace("\x00", "") or "uploaded_file"

def parse_bytes(filename, data, depth=0):
    if depth > 5:
        raise ValueError("ZIP nesting exceeds safe maximum depth of 5.")
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        return []
    if ext == ".pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        return [{"filename": safe_name(filename), "text": text}]
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [{"filename": safe_name(filename), "text": "\n".join(parts)}]
    if ext == ".txt":
        return [{"filename": safe_name(filename),
                 "text": data.decode("utf-8", errors="ignore")}]
    docs = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = [m for m in z.infolist() if not m.is_dir()]
        if len(members) > settings.MAX_ZIP_FILES:
            raise ValueError(f"ZIP contains more than {settings.MAX_ZIP_FILES} files.")
        for member in members:
            name = safe_name(member.filename)
            if Path(name).suffix.lower() in ALLOWED:
                docs.extend(parse_bytes(name, z.read(member), depth + 1))
    return docs

def parse_uploaded_files(files):
    docs = []
    for file in files:
        data = file.read()
        if not data:
            continue
        docs.extend(parse_bytes(safe_name(file.filename), data))
    for doc in docs:
        doc["text"] = " ".join(doc["text"].split())[:settings.MAX_TEXT_CHARS]
    docs = [d for d in docs if d["text"].strip()]
    if not docs:
        raise ValueError("No readable PDF, DOCX, TXT or ZIP resume was found.")
    return docs
